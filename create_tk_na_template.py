#!/usr/bin/env python3
"""
Script untuk membuat template Excel untuk Upload TK NA (Tenaga Kerja Nonaktif)
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_tk_na_template():
    """Membuat template Excel untuk upload TK NA"""
    
    # Data contoh
    data = {
        'NIK': [
            '3175051405900001',
            '3175051405900002', 
            '3175051405900003'
        ],
        'Nama': [
            'John Doe',
            'Jane Smith',
            'Ahmad Rahman'
        ],
        'Alasan Nonaktif': [
            'Resign',
            'Meninggal',
            'Kontrak Berakhir'
        ]
    }
    
    # Buat DataFrame
    df = pd.DataFrame(data)
    
    # Simpan ke Excel dengan styling
    with pd.ExcelWriter('template_tk_na.xlsx', engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='TK_NA', index=False)
        
        # Ambil worksheet
        worksheet = writer.sheets['TK_NA']
        
        # Styling header
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        # Border
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Apply styling ke header
        for col in range(1, 4):  # Kolom A, B, C
            cell = worksheet.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Apply border ke semua cell
        for row in range(1, 5):  # 4 baris (header + 3 data)
            for col in range(1, 4):
                worksheet.cell(row=row, column=col).border = thin_border
        
        # Set column width
        worksheet.column_dimensions['A'].width = 20  # NIK
        worksheet.column_dimensions['B'].width = 25  # Nama
        worksheet.column_dimensions['C'].width = 20  # Alasan Nonaktif
        
        # Freeze header row
        worksheet.freeze_panes = 'A2'
    
    print("Template TK NA berhasil dibuat: template_tk_na.xlsx")
    print("\nFormat Template:")
    print("   Kolom A: NIK (16 digit)")
    print("   Kolom B: Nama Lengkap")
    print("   Kolom C: Alasan Nonaktif")
    print("\nAlasan Nonaktif yang umum:")
    print("   - Resign")
    print("   - Meninggal")
    print("   - Kontrak Berakhir")
    print("   - PHK")
    print("   - Pensiun")

if __name__ == "__main__":
    create_tk_na_template()
