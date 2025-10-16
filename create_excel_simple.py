import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Data test dengan format yang benar sesuai SIPP
data = [
    ['PEG001', 'John Doe', 'S.Kom', '021', '7654321', '', '021', '021-1234', '1234', '081234567890', 'john.doe@example.com', 'Jakarta', '1990-05-14', 'Siti Aminah', 'KTP', '3175051405900001', '2016-05-14', 'Laki-laki', 'Jalan Mandaka No. 10', '2020-07-01', 'Belum Kawin', 'O', '012345678901234', 'ID', '5000000', 'Jl. Merpati 12 Jakarta', '10130', 'Kantor Pusat', 'PKWTT', '2020-07-01', '', '0', 'WNI'],
    ['PEG002', 'Jane Smith', 'S.E', '022', '8765432', '', '022', '022-5678', '5678', '082345678901', 'jane.smith@example.com', 'Bandung', '1988-03-22', 'Maria Sari', 'KTP', '3175051405900002', '2018-03-22', 'Perempuan', 'Jalan Sudirman No. 5', '2021-01-15', 'Kawin', 'A', '012345678901235', 'ID', '6000000', 'Jl. Gatot Subroto 25 Bandung', '40112', 'Cabang Bandung', 'PKWTT', '2021-01-15', '', '0', 'WNI'],
    ['PEG003', 'Ahmad Rahman', 'S.T', '024', '9876543', '', '024', '024-9012', '9012', '083456789012', 'ahmad.rahman@example.com', 'Semarang', '1992-11-08', 'Fatimah Zahra', 'KTP', '3175051405900003', '2020-11-08', 'Laki-laki', 'Jalan Diponegoro No. 15', '2022-03-10', 'Belum Kawin', 'B', '012345678901236', 'ID', '4500000', 'Jl. Imam Bonjol 8 Semarang', '50111', 'Cabang Semarang', 'PKWT', '2022-03-10', '2024-03-10', '0', 'WNI']
]

# Header sesuai template SIPP
headers = ['NO_PEGAWAI', 'NAMA_LENGKAP', 'GELAR', 'TELEPON_AREA_RUMAH', 'TELEPON_RUMAH', '', 'TELEPON_AREA_KANTOR', 'TELEPON_KANTOR', 'TELEPON_EXT_KANTOR', 'HP', 'EMAIL', 'TEMPAT_LAHIR', 'TANGGAL_LAHIR', 'NAMA_IBU_KANDUNG', 'JENIS_IDENTITAS', 'NO_IDENTITAS', 'MASA_LAKU_IDENTITAS', 'JENIS_KELAMIN', 'SURAT_MENYURAT_KE', 'TANGGAL_KEPESERTAAN', 'STATUS_KAWIN', 'GOLONGAN_DARAH', 'NPWP', 'KODE_NEGARA', 'UPAH', 'ALAMAT', 'KODE_POS', 'LOKASI_PEKERJAAN', 'STATUS_PEGAWAI', 'TGL_AWAL_BEKERJA', 'TGL_AKHIR_KONTRAK', 'RAPEL', 'NATIONALITY']

# Buat workbook baru
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"

# Tulis header
for col, header in enumerate(headers, 1):
    ws.cell(row=1, column=col, value=header)

# Tulis data
for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        ws.cell(row=row_idx, column=col_idx, value=value)

# Format header
header_font = Font(bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center')

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment

# Set column widths
column_widths = [12, 20, 10, 15, 15, 5, 15, 15, 15, 15, 25, 15, 15, 20, 15, 20, 15, 15, 25, 15, 15, 15, 20, 10, 15, 25, 10, 20, 15, 15, 15, 10, 10]

for i, width in enumerate(column_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

# Simpan file
filename = 'template_sipp_valid.xlsx'
wb.save(filename)

print(f"✅ Template SIPP Excel berhasil dibuat: {filename}")
print(f"📊 Total data: {len(data)}")
print(f"📋 Kolom: {len(headers)}")
print(f"📝 Format: Sesuai spesifikasi SIPP BPJamsostek")







